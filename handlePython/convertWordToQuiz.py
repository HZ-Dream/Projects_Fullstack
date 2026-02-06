import json
import os
from docx import Document
from google import genai  # Import đúng package mới

# API key từ env (khuyến nghị), tên env phải là GEMINI_API_KEY
# Nếu không có env, pass api_key="your_key_here" khi tạo Client
client = genai.Client()

def read_word_text(file_path):
    if not os.path.exists(file_path):
        print(json.dumps({"error": "File không tồn tại"}, ensure_ascii=False))
        return ""
    
    doc = Document(file_path)
    content = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                content.append(" | ".join(row_text))

    return "\n".join(content)

def smart_truncate(text, max_chars=15000):  # Tăng lên 15k chars để an toàn hơn (~3000-4000 tokens)
    if len(text) > max_chars:
        print("Warning: Text truncated for safety.")
        return text[:max_chars] + "... (truncated)"
    return text

def generate_quiz(text):
    prompt = f"""
Bạn là chuyên gia soạn đề thi trắc nghiệm theo yêu cầu sau:
- Tên tiêu đề bài quiz ngắn gọn, bao quát nội dung tài liệu
- Số câu hỏi tuỳ theo lượng thông tin của tài liệu
- Mỗi câu có 4 đáp án
- Chỉ 1 đáp án đúng
- KHÔNG giải thích gì thêm 
- Tạo JSON đúng định dạng sau:

{{
  "title": "Tiêu đề ngắn gọn bao quát nội dung",
  "quiz": [
    {{
      "questionText": "Câu hỏi?",
      "options": ["Câu trả lời", "Câu trả lời", "Câu trả lời", "Câu trả lời"],
      "correctAnswers": ["Đáp án"]
    }}
  ]
}}

Nội dung tài liệu:
{text}
"""

    try:
        response = client.models.generate_content(
            model="gemini-2.5-flash", 
            contents=prompt,
            config={
                "temperature": 0.3,
                "response_mime_type": "application/json"
            }
        )
        
        # response.text là JSON string sạch
        return json.loads(response.text)
    
    except Exception as e:
        print(json.dumps({"error": f"Lỗi Gemini: {str(e)}"}, ensure_ascii=False))
        return None

def word_to_quiz(word_file):
    raw_text = read_word_text(word_file)
    
    if not raw_text.strip():
        print(json.dumps({"error": "File không có nội dung văn bản"}, ensure_ascii=False))
        return

    clean_text = smart_truncate(raw_text)

    quiz_data = generate_quiz(clean_text)
    
    if quiz_data:
        print(json.dumps(quiz_data, ensure_ascii=False))  # Chỉ in JSON sạch
    else:
        print(json.dumps({"error": "Không generate được quiz từ Gemini"}, ensure_ascii=False))

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Không có đường dẫn file"}, ensure_ascii=False))
        sys.exit(1)
    
    word_to_quiz(sys.argv[1])